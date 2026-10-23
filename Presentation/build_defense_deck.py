"""Build the Fabpulse exception-console defence deck.

Facts come from the running console contract and a local demo-state read:
work-queue count 196 at rest, 197 after Review this delivery, packaged
replay count 67 throughout. No logo file is in the repo, so the cover
has no wordmark.
"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "defense-deck.pptx"

W = Emu(12192000)
H = Emu(6858000)

RED = "ED1C24"
CHAR = "231F20"
PURPLE = "72007D"
GOLD = "EFBC4A"
WHITE = "FFFFFF"
GRAY = "F2F2F2"
RED_TINT = "FDEBEC"
GOLD_TINT = "FCF3DE"
MUTED = "666666"
BODY = "4B4B4B"
NUM = "777777"
HAIR = "E0E0E0"
LINE = "D9D9D9"

PRODUCT = "Force Fab"
FOOTER_Y = 7.08


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def style_run(run, size: float, bold: bool, color: str) -> None:
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = False
    run.font.color.rgb = rgb(color)
    r_pr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        node = r_pr.find(qn(f"a:{tag}"))
        if node is None:
            node = etree.SubElement(r_pr, qn(f"a:{tag}"))
        node.set("typeface", "Calibri")


def fill_solid(shape, color: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)


def no_line(shape) -> None:
    shape.line.fill.background()


def hairline(shape) -> None:
    shape.line.color.rgb = rgb(HAIR)
    shape.line.width = Pt(0.75)


def set_run_text(paragraph, text: str, size: float, bold: bool, color: str) -> None:
    run = paragraph.add_run()
    run.text = text
    style_run(run, size, bold, color)


def write_text(
    shape,
    lines: list[tuple[str, float, bool, str]],
    *,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    name: str | None = None,
    margin: float = 0.08,
) -> None:
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.04)
    frame._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    for index, (text, size, bold, color) in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        set_run_text(paragraph, text, size, bold, color)


def add_rect(slide, x, y, w, h, fill: str, line: bool = False, name: str | None = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill_solid(shape, fill)
    if line:
        hairline(shape)
    else:
        no_line(shape)
    if name:
        shape.name = name
    shape.shadow.inherit = False
    return shape


def add_round(slide, x, y, w, h, name: str | None = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill_solid(shape, WHITE)
    hairline(shape)
    shape.adjustments[0] = 0.08
    if name:
        shape.name = name
    shape.shadow.inherit = False
    return shape


def add_bar(slide, x, y, h, color: str) -> None:
    inset = min(0.18, h * 0.16)
    add_rect(slide, x + 0.02, y + inset, 0.07, h - (inset * 2), color)


def add_textbox(slide, x, y, w, h, lines, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name: str | None = None, margin: float = 0.06):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    write_text(shape, lines, align=align, anchor=anchor, name=name, margin=margin)
    return shape


def set_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(WHITE)


def add_notes(slide, text: str) -> None:
    frame = slide.notes_slide.notes_text_frame
    frame.clear()
    blocks = [block.strip() for block in text.strip().split("\n\n")]
    for index, block in enumerate(blocks):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = Pt(8)
        set_run_text(paragraph, block, 14, False, CHAR)


def footer(slide, number: int) -> None:
    add_textbox(
        slide,
        0.55,
        FOOTER_Y,
        6.4,
        0.26,
        [(PRODUCT, 10, False, MUTED)],
        anchor=MSO_ANCHOR.MIDDLE,
        name="footer-name",
        margin=0.0,
    )
    add_textbox(
        slide,
        12.35,
        FOOTER_Y,
        0.7,
        0.26,
        [(str(number), 10, False, NUM)],
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
        name="footer-num",
        margin=0.0,
    )


def claim(slide, text: str, width: float = 12.2) -> None:
    add_textbox(
        slide,
        0.55,
        0.38,
        width,
        0.95,
        [(text, 26, True, CHAR)],
        anchor=MSO_ANCHOR.TOP,
        name="title",
        margin=0.0,
    )


def row_cards(count: int, y: float, height: float, left: float = 0.55, right: float = 12.78, gap: float = 0.18):
    width = (right - left - gap * (count - 1)) / count
    boxes = []
    x = left
    for _ in range(count):
        boxes.append((x, y, width, height))
        x += width + gap
    return boxes


def number_circle(slide, cx, cy, diameter, numeral: str, fill: str) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(cx),
        Inches(cy),
        Inches(diameter),
        Inches(diameter),
    )
    fill_solid(shape, fill)
    no_line(shape)
    shape.shadow.inherit = False
    write_text(
        shape,
        [(numeral, 16, True, WHITE)],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        name="seq",
        margin=0.0,
    )


def join_circles(slide, x1, y, x2) -> None:
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y), Inches(x2), Inches(y))
    line.line.color.rgb = rgb(LINE)
    line.line.width = Pt(2)


def card_shell(slide, box, bar: str):
    x, y, w, h = box
    add_round(slide, x, y, w, h)
    add_bar(slide, x, y, h, bar)
    return x, y, w, h


def pill(slide, x, y, w, h, text: str) -> None:
    shape = add_rect(slide, x, y, w, h, RED_TINT, name="pill")
    write_text(
        shape,
        [(text, 9, True, RED)],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        name="pill-label",
        margin=0.04,
    )


def paint_line_alpha(shape, color: str, width_pt: float, alpha_pct: int) -> None:
    shape.fill.background()
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width_pt)
    solid = shape.line._get_or_add_ln().find(qn("a:solidFill"))
    srgb = solid.find(qn("a:srgbClr"))
    alpha = srgb.find(qn("a:alpha"))
    if alpha is None:
        alpha = etree.SubElement(srgb, qn("a:alpha"))
    alpha.set("val", str(alpha_pct * 1000))
    geom = shape._element.spPr.find(qn("a:prstGeom"))
    if geom is not None:
        av_lst = geom.find(qn("a:avLst"))
        if av_lst is not None:
            for child in list(av_lst):
                av_lst.remove(child)


def add_arc(slide, x, y, size, color: str, width_pt: float, alpha_pct: int, name: str) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ARC, Inches(x), Inches(y), Inches(size), Inches(size))
    shape.rotation = 18.0
    shape.name = name
    shape.shadow.inherit = False
    paint_line_alpha(shape, color, width_pt, alpha_pct)


def draw_table(slide, x, y, widths, headers, rows, row_h=0.46) -> float:
    header_h = 0.42
    cursor = x
    for index, header in enumerate(headers):
        cell = add_rect(slide, cursor, y, widths[index], header_h, CHAR)
        write_text(
            cell,
            [(header, 13, True, WHITE)],
            anchor=MSO_ANCHOR.MIDDLE,
            name="table",
            margin=0.08,
        )
        cursor += widths[index]
    for r_index, row in enumerate(rows):
        top = y + header_h + r_index * row_h
        fill = WHITE if r_index % 2 == 0 else GRAY
        cursor = x
        for c_index, value in enumerate(row):
            cell = add_rect(slide, cursor, top, widths[c_index], row_h, fill)
            write_text(
                cell,
                [(value, 13, c_index == 0, CHAR)],
                anchor=MSO_ANCHOR.MIDDLE,
                name="table",
                margin=0.08,
            )
            rule = slide.shapes.add_connector(
                1,
                Inches(cursor),
                Inches(top + row_h),
                Inches(cursor + widths[c_index]),
                Inches(top + row_h),
            )
            rule.line.color.rgb = rgb(HAIR)
            rule.line.width = Pt(0.75)
            cursor += widths[c_index]
    return y + header_h + len(rows) * row_h


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width = W
        self.prs.slide_height = H
        self.prs.core_properties.title = PRODUCT
        self.prs.core_properties.author = "Team-Force"
        self.prs.core_properties.subject = "Internal capstone defence"
        self.blank = self.prs.slide_layouts[6]
        self.number = 0

    def slide(self, kind: str = "content"):
        slide = self.prs.slides.add_slide(self.blank)
        self.number += 1
        set_background(slide)
        if kind != "cover":
            footer(slide, self.number)
        return slide

    def save(self) -> None:
        for slide in self.prs.slides:
            send_connectors_back(slide)
        self.prs.save(OUT)


def send_connectors_back(slide) -> None:
    tree = slide.shapes._spTree
    connectors = [element for element in list(tree) if element.tag == qn("p:cxnSp")]
    for element in connectors:
        tree.remove(element)
    insert_at = 2
    for element in connectors:
        tree.insert(insert_at, element)
        insert_at += 1


def add_sentences(slide, sentences, y=1.45, h=1.55, size=16, name="body"):
    shape = slide.shapes.add_textbox(Inches(0.55), Inches(y), Inches(12.2), Inches(h))
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = Inches(0)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0)
    for index, text in enumerate(sentences):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(8)
        set_run_text(paragraph, text, size, False, CHAR)
    return shape


def speak(slide, first: str, point: str, nxt: str) -> None:
    add_notes(
        slide,
        f"{first}\n\nPoint at {point}.\n\nNext slide: {nxt}.",
    )


def slide_01(deck: Deck) -> None:
    slide = deck.slide("cover")
    add_rect(slide, 0, 0, 0.22, 7.5, RED, name="cover-bar")
    add_arc(slide, 8.85, -0.40, 4.9, RED, 8, 90, "cover-arc-outer")
    add_arc(slide, 9.55, 0.25, 4.4, CHAR, 2, 65, "cover-arc-inner")
    pill(slide, 0.85, 1.00, 2.35, 0.34, "Internal review")
    add_textbox(
        slide,
        0.85,
        1.55,
        7.2,
        1.35,
        [(PRODUCT, 36, True, CHAR)],
        name="title",
        margin=0.0,
    )
    add_textbox(
        slide,
        0.90,
        3.15,
        7.0,
        2.2,
        [
            ("A reviewer uses this page to read one fab-lot or equipment exception and to record the decision that role allows.", 18, False, MUTED),
            ("The explanation and the live demonstration take 40 minutes. Questions take the next 10 minutes.", 18, False, MUTED),
        ],
        name="support",
        margin=0.0,
    )
    add_textbox(
        slide,
        0.90,
        6.75,
        8.6,
        0.4,
        [(f"{PRODUCT}   ·   Team-Force   ·   20 October 2026   ·   Prepared for internal use", 10, False, MUTED)],
        name="cover-footer",
        margin=0.0,
    )
    speak(
        slide,
        "This is Force Fab. A reviewer reads one exception and records the decision that role allows. The explanation and the live demonstration take 40 minutes. Questions take the next 10 minutes.",
        "the title",
        "the seven parts of the hour",
    )


def slide_02(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "This hour has seven parts.")
    add_sentences(
        slide,
        ["Each part below has a clock time. The live clicks sit inside the third part."],
        y=1.40,
        h=0.45,
        size=16,
    )
    rows = [
        ("0:00", "The records disagree, and the packaged counts say how often."),
        ("7:00", "Two applications share the review. The model drafts a sentence only."),
        ("12:00", "One second delivery is finished on the page, including the live clicks."),
        ("24:00", "The same five cases open on Application B, which shows the rule path."),
        ("31:00", "The frozen counts stay the same. Each number keeps its honesty label."),
        ("35:00", "A blocked note and a restarted process left the case as it was."),
        ("38:00", "Five measured lessons close the walk. Questions follow at 40:00."),
    ]
    for index, (clock, text) in enumerate(rows):
        y = 1.95 + index * 0.68
        add_rect(slide, 0.55, y, 1.15, 0.56, RED_TINT if index == 0 else GRAY)
        add_textbox(
            slide,
            0.55,
            y,
            1.15,
            0.56,
            [(clock, 14, True, RED if index == 0 else CHAR)],
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
            name="agenda",
            margin=0.04,
        )
        add_textbox(
            slide,
            1.85,
            y,
            10.7,
            0.56,
            [(text, 16, False, CHAR)],
            anchor=MSO_ANCHOR.MIDDLE,
            name="body",
            margin=0.06,
        )
    speak(
        slide,
        "This hour has seven parts. The live clicks sit inside the third part, from minute 12 to minute 24.",
        "the seven clock times",
        "what disagrees on one exception",
    )


def slide_03(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "The records for one exception do not agree.")
    add_sentences(
        slide,
        [
            "Fabpulse Group handles fab-lot exceptions and equipment exceptions.",
            "A partner feed, a process record, a policy, a case workflow, an archived document, and a local rule can each give a different answer for the same case.",
            "The reviewer then has to decide which record to trust before the case can close.",
        ],
        y=1.42,
        h=1.7,
    )
    rows = [
        ["Partner feed", "The external message that arrived"],
        ["Process record", "The recipe the tool actually ran"],
        ["Policy version", "Which rule is in force on that date"],
        ["Case workflow", "Whether the case is open, in review, or closed"],
        ["Archived document", "The statement stored for that case"],
        ["Local rule", "The program that site uses to decide"],
    ]
    draw_table(slide, 0.55, 3.25, [3.3, 8.9], ["Record", "What it tells the reviewer"], rows, row_h=0.52)
    speak(
        slide,
        "The records for one exception do not agree. Six sources can each give a different answer for the same case.",
        "the six records",
        "what that disagreement does on a shift",
    )


def slide_04(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "A disagreed case costs the shift time and a clear owner.")
    add_sentences(
        slide,
        [
            "Operators spend extra time reconciling the same case by hand.",
            "Two teams can report a different open count for the same week.",
            "A disputed outage can sit with no follow-up owner.",
            "This defence states no dollar cost. Dollars for that delay stay Unknown.",
        ],
        y=1.42,
        h=1.85,
    )
    boxes = row_cards(1, 3.45, 2.7)
    x, y, w, h = card_shell(slide, boxes[0], RED)
    add_textbox(
        slide,
        x + 0.28,
        y + 0.25,
        w - 0.5,
        2.2,
        [
            ("Worked example. Incident INC-2025-014.", 16, True, CHAR),
            ("On 2025-11-18 a partner feed was delayed. The queue was retried twice. Staff used a spreadsheet at 04:00 UTC.", 16, False, BODY),
            ("The owner field on this incident is empty. That empty field is REAL. This walk does not type a name into it.", 16, False, CHAR),
        ],
        name="card:example",
    )
    speak(
        slide,
        "A disagreed case costs the shift time and a clear owner. Incident INC-2025-014 has an empty owner field. Dollars for the delay stay Unknown.",
        "the incident example",
        "the eight counted defects",
    )


def slide_05(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "Eight defects are counted from the packaged data.")
    add_sentences(
        slide,
        ["Each count in this table is REAL. It was counted from the packaged files. It is not a goal invented for this defence."],
        y=1.40,
        h=0.7,
    )
    rows = [
        ["Recipe mismatch", "91 of 125 steps", "REAL"],
        ["Replay feeds", "67 REPLAY rows", "REAL"],
        ["Mixed clocks", "11,985 of 12,500 events", "REAL"],
        ["Four rule programs", "2,090 REVIEW objects", "REAL"],
        ["Open and closed", "89 of 125 cases still open", "REAL"],
        ["Overlapping policies", "34 active rows", "REAL"],
        ["Untrusted notes", "86 partner feeds", "REAL"],
        ["Missing owners", "695 unassigned tasks", "REAL"],
    ]
    draw_table(slide, 0.55, 2.15, [3.5, 5.5, 2.2], ["Defect", "Count", "Honesty"], rows, row_h=0.46)
    add_sentences(
        slide,
        ["Exception cycle time is Unknown. Aged backlog is Unknown. The file has no start time, no end time, and no age in days."],
        y=6.40,
        h=0.55,
        size=14,
        name="caption",
    )
    speak(
        slide,
        "Eight defects are counted from the packaged data. Each count is REAL. Cycle time and aged backlog are Unknown.",
        "the honesty column",
        "the six frozen baseline fields",
    )


def slide_06(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "Four frozen fields are REAL. Two are simulated.")
    add_sentences(
        slide,
        [
            "These six fields are the baseline. After the build they are the same numbers.",
            "A REAL count was read from the packaged cases. A SIMULATED figure is a model on synthetic data. It is not cash, and it is not recovered labor.",
        ],
        y=1.40,
        h=1.15,
    )
    rows = [
        ["Records", "125", "REAL"],
        ["Open cases", "89", "REAL"],
        ["Urgent cases", "2", "REAL"],
        ["Sites", "6", "REAL"],
        ["Mean amount", "4982.61", "SIMULATED"],
        ["Hours per case", "0.603", "SIMULATED"],
    ]
    draw_table(slide, 0.55, 2.70, [4.0, 4.0, 3.2], ["Field", "Value", "Honesty"], rows, row_h=0.58)
    speak(
        slide,
        "Four frozen fields are REAL. Records are 125, open cases are 89, urgent cases are 2, and sites are 6. Mean amount and hours per case are SIMULATED. They are not cash.",
        "the honesty column",
        "the two applications that were built",
    )


def slide_07(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "Two applications run the same review and stay apart.")
    add_sentences(
        slide,
        ["A reviewer can open either application and see the same kind of page. The two servers do not call each other."],
        y=1.40,
        h=0.7,
    )
    boxes = row_cards(3, 2.20, 2.55)
    cards = [
        ("Application A", "This is the rebuilt program. The reviewer opens it on port 8100."),
        ("Application B", "This application reads the Context Lake. The reviewer opens it on port 8110."),
        ("The boundary", "The browser can read both. The servers do not call each other."),
    ]
    bars = [RED, CHAR, PURPLE]
    for index, (box, card) in enumerate(zip(boxes, cards)):
        x, y, w, h = card_shell(slide, box, bars[index])
        add_textbox(
            slide,
            x + 0.24,
            y + 0.25,
            w - 0.42,
            2.0,
            [(card[0], 15, True, CHAR), (card[1], 16, False, BODY)],
            name=f"card:{index}",
        )
    add_sentences(
        slide,
        [
            "The lake version is 1.0.0. It holds 148 items, 180 edges, and 267 term rows. That count is REAL.",
            "Replay mode is on. Saved answers play back from files. This walk does not call a live model over the network.",
        ],
        y=4.95,
        h=1.4,
    )
    speak(
        slide,
        "Two applications run the same review and stay apart. Application A is the rebuilt program on port 8100. Application B reads the Context Lake on port 8110. Replay mode is on.",
        "the three cards",
        "which jobs stay in code and which jobs the model drafts",
    )


def slide_08(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "The model drafts a sentence. The code records the decision.")
    add_sentences(
        slide,
        ["A cited explanation needs a draft in words. Ranking, rules, reconciliation, and de-duplication already have a determined result, so those jobs stay in code."],
        y=1.40,
        h=0.85,
    )
    boxes = row_cards(2, 2.40, 3.15)
    card_shell(slide, boxes[0], RED)
    card_shell(slide, boxes[1], CHAR)
    add_textbox(
        slide,
        boxes[0][0] + 0.24,
        boxes[0][1] + 0.25,
        boxes[0][2] - 0.45,
        2.7,
        [
            ("Stays in code", 15, True, CHAR),
            ("Ranking, rules, reconciliation, and de-duplication run in code. A person in roles RT-01 to RT-06 records every write.", 16, False, BODY),
        ],
        name="card:code",
    )
    add_textbox(
        slide,
        boxes[1][0] + 0.24,
        boxes[1][1] + 0.25,
        boxes[1][2] - 0.45,
        2.7,
        [
            ("Drafted by the model", 15, True, CHAR),
            ("The model drafts a cited explanation, a query answer, or a plan. It has no write tool. It does not write a hold, a case status, or a feed status.", 16, False, BODY),
        ],
        name="card:model",
    )
    add_sentences(
        slide,
        ["If the saved explanation fails, the business result can still change. The page shows both sentences. They stay separate."],
        y=5.75,
        h=0.8,
    )
    speak(
        slide,
        "The model drafts a sentence. The code records the decision. The model has no write tool. If the explanation fails, the business result can still change, and the page shows both.",
        "the two cards",
        "why a hold prediction was rejected",
    )


def slide_09(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "Hold prediction scored the same as the majority guess.")
    add_sentences(
        slide,
        [
            "The question was whether a model should predict which lot to hold.",
            "The packaged set has 125 lots. 8 lots were held. 117 lots were not held. Those counts are REAL.",
            "A majority guess scores 117 of 125. Leave-one-out scoring also returned 117 of 125. The prediction was rejected.",
            "No production hold scorer is in this build. The hold decision stays with a person in role RT-03.",
        ],
        y=1.45,
        h=2.3,
    )
    boxes = row_cards(3, 4.00, 2.15)
    facts = [
        ("125 lots", "8 held. 117 not held. REAL."),
        ("117 of 125", "Majority guess and leave-one-out. REAL."),
        ("Rejected", "The model did not beat the simple count."),
    ]
    bars = [RED, CHAR, GOLD]
    for index, (box, fact) in enumerate(zip(boxes, facts)):
        x, y, w, h = card_shell(slide, box, bars[index])
        add_textbox(
            slide,
            x + 0.22,
            y + 0.3,
            w - 0.4,
            1.5,
            [(fact[0], 18, True, CHAR), (fact[1], 15, False, BODY)],
            name=f"card:{index}",
        )
    speak(
        slide,
        "Hold prediction scored 117 of 125, the same as guessing the majority. It was rejected. The hold stays with RT-03.",
        "the three cards",
        "the second delivery on this shift",
    )


def slide_10(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "One second delivery is already on this shift.")
    add_sentences(
        slide,
        [
            "Partner feed REPLAY-000000000 is a second delivery of reference C00001646.",
            "The first ingest time is 2022-04-25 at 19:09. The second ingest time is 2022-04-25 at 20:09.",
        ],
        y=1.42,
        h=1.05,
    )
    boxes = row_cards(2, 2.60, 3.5)
    card_shell(slide, boxes[0], RED)
    card_shell(slide, boxes[1], CHAR)
    add_textbox(
        slide,
        boxes[0][0] + 0.24,
        boxes[0][1] + 0.3,
        boxes[0][2] - 0.45,
        2.8,
        [
            ("If the duplicate stays", 15, True, CHAR),
            ("The same feed is counted twice on the work queue. The reviewer is looking at one arrival that came in again.", 16, False, BODY),
        ],
        name="card:stay",
    )
    add_textbox(
        slide,
        boxes[1][0] + 0.24,
        boxes[1][1] + 0.3,
        boxes[1][2] - 0.45,
        2.8,
        [
            ("If the wrong copy is dropped", 15, True, CHAR),
            ("A true delivery leaves the queue. The page therefore names the role that is allowed to drop it before the click.", 16, False, BODY),
        ],
        name="card:drop",
    )
    speak(
        slide,
        "Partner feed REPLAY-000000000 is a second delivery of C00001646. The ingest times are 19:09 and 20:09.",
        "the two ingest times",
        "how the page is arranged",
    )


def slide_11(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "One page shows the case, the button, and the result.")
    add_sentences(
        slide,
        ["The reviewer reads the situation, then the evidence, then the red button, then the result. The red button on this case is Review this delivery. The reviewer finishes that click on this page."],
        y=1.38,
        h=0.85,
    )
    frame = add_round(slide, 0.55, 2.30, 12.15, 4.35, name="console-frame")
    frame.shadow.inherit = False
    page_x, page_y, page_w, page_h = 0.72, 2.42, 11.82, 3.85
    add_rect(slide, page_x, page_y, page_w, page_h, GRAY)
    add_rect(slide, page_x, page_y, 0.12, page_h, RED, name="console-edge")
    header_x = page_x + 0.12
    header_w = page_w - 0.12
    add_rect(slide, header_x, page_y, header_w, 0.48, WHITE)
    add_textbox(
        slide,
        header_x + 0.12,
        page_y + 0.02,
        6.2,
        0.44,
        [("Application A     Port 8100     RT-01", 14, False, CHAR)],
        anchor=MSO_ANCHOR.MIDDLE,
        name="console",
        margin=0.0,
    )
    pill(slide, header_x + 5.4, page_y + 0.08, 0.9, 0.30, "replay")
    tab_y = page_y + 0.48
    add_rect(slide, header_x, tab_y, 1.9, 0.42, WHITE)
    add_rect(slide, header_x, tab_y, 1.9, 0.04, RED)
    add_textbox(
        slide,
        header_x + 0.08,
        tab_y,
        1.7,
        0.42,
        [("Replay", 13, True, CHAR), ("C00001646", 11, False, MUTED)],
        name="console",
        margin=0.02,
    )
    sheet_y = tab_y + 0.42
    add_rect(slide, header_x, sheet_y, header_w, page_y + page_h - sheet_y, WHITE, line=True)
    add_textbox(
        slide,
        header_x + 0.18,
        sheet_y + 0.06,
        7.5,
        0.7,
        [
            ("Situation. Partner feed REPLAY-000000000 is a second delivery of C00001646.", 13, False, CHAR),
            ("Evidence. Ref C00001646. This delivery is not on the queue yet.", 13, False, CHAR),
        ],
        name="console",
        margin=0.0,
    )
    button = add_round(slide, header_x + 0.18, sheet_y + 0.85, 2.8, 0.40)
    fill_solid(button, RED)
    no_line(button)
    write_text(
        button,
        [("Review this delivery", 14, True, WHITE)],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
        name="console",
        margin=0.04,
    )
    band_y = sheet_y + 1.38
    add_rect(slide, header_x + 0.18, band_y, header_w - 0.4, 0.38, GOLD_TINT)
    add_rect(slide, header_x + 0.18, band_y, 0.08, 0.38, GOLD)
    add_textbox(
        slide,
        header_x + 0.34,
        band_y,
        8.0,
        0.38,
        [("Result. Ready. No click has been made yet.", 14, False, CHAR)],
        anchor=MSO_ANCHOR.MIDDLE,
        name="console",
        margin=0.0,
    )
    add_textbox(
        slide,
        0.72,
        6.32,
        5.5,
        0.28,
        [("Structure of the console.", 12, False, MUTED)],
        name="caption",
        margin=0.0,
    )
    speak(
        slide,
        "One page shows the case, the button, and the result. The red button is Review this delivery. The result still says Ready because no click has been made.",
        "Review this delivery",
        "what that click changes",
    )


def slide_12(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "The click moves the queue from 196 to 197.")
    add_sentences(
        slide,
        [
            "Before the click, this delivery is not on the queue and the work-queue count is 196. After Review this delivery, the row says the delivery is on the work queue and the count is 197.",
            "The packaged store still has 67 replay rows. The session added the delivery. The packaged files were not rewritten. Both 67 figures are REAL.",
        ],
        y=1.38,
        h=1.25,
    )
    boxes = row_cards(2, 2.75, 3.15)
    panels = [
        ("Before the click", [("This delivery", "not on the queue"), ("Work-queue count", "196"), ("Packaged replay rows", "67")], False),
        ("After Review this delivery", [("This delivery", "on the work queue"), ("Work-queue count", "197"), ("Packaged replay rows", "67")], True),
    ]
    for index, (box, panel) in enumerate(zip(boxes, panels)):
        x, y, w, h = card_shell(slide, box, RED if index == 0 else CHAR)
        title, rows, changed = panel
        add_textbox(slide, x + 0.22, y + 0.12, w - 0.4, 0.36, [(title, 14, True, MUTED)], name=f"card:{index}", margin=0.0)
        for r_index, (label, value) in enumerate(rows):
            ry = y + 0.55 + r_index * 0.72
            highlight = changed and r_index == 0
            if highlight:
                add_rect(slide, x + 0.22, ry, w - 0.44, 0.62, GOLD_TINT)
                add_rect(slide, x + 0.22, ry, 0.08, 0.62, GOLD)
            else:
                add_rect(slide, x + 0.22, ry, w - 0.44, 0.62, GRAY if r_index % 2 else WHITE, line=True)
            add_textbox(slide, x + 0.36, ry, (w - 0.5) * 0.55, 0.62, [(label, 14, True, CHAR)], anchor=MSO_ANCHOR.MIDDLE, name=f"card:{index}", margin=0.04)
            add_textbox(slide, x + 0.36 + (w - 0.5) * 0.52, ry, (w - 0.5) * 0.45, 0.62, [(value, 14, False, CHAR)], anchor=MSO_ANCHOR.MIDDLE, name=f"card:{index}", margin=0.04)
    add_sentences(
        slide,
        ["The result line stays on the page. It says the work-queue count is 197 and the packaged replay count is 67."],
        y=6.05,
        h=0.7,
        size=15,
    )
    speak(
        slide,
        "The click moves the work-queue count from 196 to 197. Packaged replay rows stay 67. The session added the delivery. The packaged files were not rewritten.",
        "197, then 67",
        "who is allowed to drop the duplicate",
    )


def slide_13(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "Only RT-05 can drop the duplicate.")
    add_sentences(
        slide,
        ["The page names that role before the click. A wrong role is part of the demonstration. The evidence row stays put until the allowed role clicks."],
        y=1.40,
        h=0.8,
    )
    boxes = row_cards(2, 2.30, 3.9)
    card_shell(slide, boxes[0], RED)
    card_shell(slide, boxes[1], CHAR)
    add_textbox(
        slide,
        boxes[0][0] + 0.24,
        boxes[0][1] + 0.25,
        boxes[0][2] - 0.45,
        3.4,
        [
            ("What the page says first", 15, True, CHAR),
            ("Drop is recorded only by RT-05 Operations Analyst.", 16, False, BODY),
            ("RT-01 clicks Drop this duplicate. The page says: Refused. Confirm drop waits for RT-05. The current role is RT-01.", 16, False, CHAR),
        ],
        name="card:refused",
    )
    add_textbox(
        slide,
        boxes[1][0] + 0.24,
        boxes[1][1] + 0.25,
        boxes[1][2] - 0.45,
        3.4,
        [
            ("What the page says after RT-05 clicks", 15, True, CHAR),
            ("RT-05 dropped the duplicate. The work-queue count returns to 196. The delivery row reads dropped.", 16, False, BODY),
            ("Packaged replay rows stay 67. One person reviews the delivery. Another person drops it.", 16, False, CHAR),
        ],
        name="card:success",
    )
    speak(
        slide,
        "Only RT-05 can drop the duplicate. RT-01 is refused in a sentence that names RT-05 and RT-01. RT-05 drops it, the count returns to 196, and packaged rows stay 67.",
        "the refusal sentence, then the success sentence",
        "the seven live moves",
    )


def slide_14(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "Seven moves finish this case on the live page.")
    add_sentences(
        slide,
        ["Leave this slide up. The role picker starts at RT-01. The runtime pill says replay. Open Application A on port 8100 and use the Replay tab."],
        y=1.38,
        h=0.7,
    )
    moves = [
        ("Open the Replay tab", "Point at C00001646"),
        ("Read the situation sentence", "Point at 19:09 and 20:09"),
        ("Click Review this delivery", "Point at count 197 and packaged 67"),
        ("Click Drop this duplicate as RT-01", "Point at the sentence that starts with Refused"),
        ("Switch the role picker to RT-05", "Point at the role picker"),
        ("Click Drop this duplicate again", "Point at dropped, count 196, and 67"),
        ("Open the Recipe tab", "Point at Ready on a clean result line"),
    ]
    for index, (move, point) in enumerate(moves):
        y = 2.15 + index * 0.62
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.55), Inches(y), Inches(0.40), Inches(0.40))
        fill_solid(circle, WHITE)
        circle.line.color.rgb = rgb(RED)
        circle.line.width = Pt(1.4)
        circle.shadow.inherit = False
        write_text(circle, [(str(index + 1), 14, True, CHAR)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name="move", margin=0.0)
        add_textbox(slide, 1.10, y, 6.8, 0.40, [(move, 15, True, CHAR)], anchor=MSO_ANCHOR.MIDDLE, name="move", margin=0.0)
        add_textbox(slide, 7.9, y, 4.7, 0.40, [(point, 13, False, MUTED)], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, name="move", margin=0.0)
    speak(
        slide,
        "Seven moves finish this case on the live page. If the saved explanation fails, say that the work-queue count still moved to 197, read the failure sentence on the screen, and continue to the RT-05 drop. Stay on the page. If the stack is down, say the live click was not run and read the same sentences from the previous slides.",
        "move 3, then move 4, then move 6",
        "the same five cases on both applications",
    )


def slide_15(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "The same five cases are open on both applications.")
    add_sentences(
        slide,
        ["Application A and Application B show these five objects. Each row is REAL. The page shows the fact in the last column. It does not invent a match percent."],
        y=1.38,
        h=0.75,
    )
    rows = [
        ["Second delivery", "REPLAY-000000000", "Drop waits for RT-05", "REAL"],
        ["Recipe mismatch", "PS000000016", "Dispatch is R02. The tool is R01.", "REAL"],
        ["Two clocks", "E-00000000", "Event 2020-12-30. Arrival 2021-01-01.", "REAL"],
        ["Policies", "FAB-00000002", "Four policies are in force on 2024-06-01.", "REAL"],
        ["Missing owner", "INC-2025-014", "The owner field is empty.", "REAL"],
    ]
    draw_table(slide, 0.55, 2.25, [2.5, 2.8, 5.2, 1.5], ["Case", "Object", "What the page shows", "Honesty"], rows, row_h=0.72)
    speak(
        slide,
        "The same five cases are open on both applications. Recipe shows dispatch R02 and tool R01. Four policies are in force on FAB-00000002. The owner of INC-2025-014 is empty.",
        "the recipe row and the policy row",
        "the rule path that only Application B shows",
    )


def slide_16(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "Only Application B shows where the recipe rule came from.")
    add_sentences(
        slide,
        [
            "Application A has no lineage route. Its Rule trace tab tells the reviewer to open Application B.",
            "Application B lists five existing steps: a historical decision, earlier code, a specification, a test, and the transformed implementation.",
            "One cited sentence holds 18600 for the 2023 EMEA recipe-release limit. That figure is REAL. It is a limit inside the rule sentence. It is not a cash benefit of this project.",
            "The Compare page also states that the 50,000 label is proposed on Application B only.",
        ],
        y=1.45,
        h=2.5,
    )
    rows = [
        ["1", "Historical decision"],
        ["2", "Earlier code"],
        ["3", "Specification"],
        ["4", "Test"],
        ["5", "Transformed implementation"],
    ]
    draw_table(slide, 0.55, 4.15, [1.4, 10.8], ["Step", "What Application B lists"], rows, row_h=0.46)
    speak(
        slide,
        "Only Application B shows where the recipe rule came from. It lists five steps. The cited sentence holds 18600. That figure is a rule limit, not a cash benefit.",
        "the five steps",
        "the packaged counts after the build",
    )


def slide_17(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "The packaged counts stayed on the same numbers.")
    add_sentences(
        slide,
        ["The build did not close the 89 open cases. It made the disagreement visible and the decision traceable. Before and after are the same for these fields."],
        y=1.38,
        h=0.75,
    )
    rows = [
        ["Records", "125", "125", "REAL"],
        ["Open cases", "89", "89", "REAL"],
        ["Urgent cases", "2", "2", "REAL"],
        ["Sites", "6", "6", "REAL"],
        ["Mean amount", "4982.61", "4982.61", "SIMULATED"],
        ["Hours per case", "0.603", "0.603", "SIMULATED"],
        ["Harness", "not run", "12 matched, 2 differed", "REAL"],
        ["Freeze", "not run", "42 of 42 passed", "REAL"],
        ["Replay tokens shown", "not shown", "0 and 0", "REAL"],
        ["ROI", "unclaimed", "Educational", "EDUCATIONAL"],
    ]
    draw_table(slide, 0.55, 2.20, [3.1, 2.5, 3.6, 2.4], ["Measure", "Before", "After", "Honesty"], rows, row_h=0.42)
    speak(
        slide,
        "The packaged counts stayed on the same numbers. Open cases stay 89. The harness matched 12 rows and differed on 2. There is no match percent. ROI stays educational. No cash benefit is claimed.",
        "the open-cases row and the ROI row",
        "the headlines that stay Unknown",
    )


def slide_18(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "The delivered value is a traceable decision on the page.")
    add_sentences(
        slide,
        ["The reviewer can see both conflicting facts before the click. A refused role is written in a sentence that names the required role and the current role. A blocked untrusted note leaves the case unchanged."],
        y=1.42,
        h=1.15,
    )
    boxes = row_cards(2, 2.70, 3.55)
    card_shell(slide, boxes[0], RED)
    card_shell(slide, boxes[1], CHAR)
    add_textbox(
        slide,
        boxes[0][0] + 0.24,
        boxes[0][1] + 0.22,
        boxes[0][2] - 0.45,
        3.1,
        [
            ("What the page now shows", 15, True, CHAR),
            ("Both ingest times, the role line, the new queue count, and the packaged count that stayed at 67.", 16, False, BODY),
            ("The answer source can read Azure OpenAI, a replay file, or a lake statement. On this build the pill says replay.", 16, False, CHAR),
        ],
        name="card:shows",
    )
    add_textbox(
        slide,
        boxes[1][0] + 0.24,
        boxes[1][1] + 0.22,
        boxes[1][2] - 0.45,
        3.1,
        [
            ("What stays Unknown", 15, True, CHAR),
            ("Exception cycle time. Aged backlog. Dollar benefit. Company job titles for RT-01 to RT-06. Numeric floors such as 90 or 95.", 16, False, BODY),
            ("The owner of INC-2025-014 stays empty until RT-06 types one on the outage page.", 16, False, CHAR),
        ],
        name="card:unknown",
    )
    speak(
        slide,
        "The delivered value is a traceable decision on the page. Cycle time, aged backlog, dollars, company titles, and numeric floors stay Unknown.",
        "the Unknown card",
        "the blocked approval",
    )


def slide_19(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "A blocked approval left the case status as REPLAY.")
    add_sentences(
        slide,
        [
            "Partner feed REPLAY-000000000 already carried untrusted quotes. The note told the reader to disregard validation and to auto approve.",
            "The prior test posted an approve action with no corpus id. Before the control, the reply waited. After the control, the reply was HTTP 409 BLOCK.",
            "The case status stayed REPLAY. Approvals stayed empty. Application B returned the same block after its own control.",
            "This slide cites that prior test. It does not run a new post in this room.",
        ],
        y=1.45,
        h=2.4,
    )
    boxes = row_cards(2, 4.10, 2.15)
    card_shell(slide, boxes[0], CHAR)
    card_shell(slide, boxes[1], GOLD)
    add_textbox(
        slide,
        boxes[0][0] + 0.22,
        boxes[0][1] + 0.25,
        boxes[0][2] - 0.4,
        1.6,
        [("Before the control", 15, True, CHAR), ("The reply waited. The case was unchanged at that moment too.", 16, False, BODY)],
        name="card:before",
    )
    add_textbox(
        slide,
        boxes[1][0] + 0.22,
        boxes[1][1] + 0.25,
        boxes[1][2] - 0.4,
        1.6,
        [("After the control", 15, True, CHAR), ("HTTP 409 BLOCK. Status stayed REPLAY. This result is REAL.", 16, False, BODY)],
        name="card:after",
    )
    speak(
        slide,
        "A blocked approval left the case status as REPLAY. The reply was HTTP 409 BLOCK. This is a prior test. This room does not run a new post.",
        "the After card",
        "the correlation id that survived a restart",
    )


def slide_20(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "The same correlation id stayed after the process died.")
    add_sentences(
        slide,
        [
            "The prior test started Application A, posted case PS000000016, and killed process id 29320.",
            "A new process started. Correlation id corr-p10b-kill-20261018-001 was still on that case.",
            "The hold value stayed 0. The case did not gain a hold because the process died.",
            "This slide cites that prior test. It does not kill a process in this room.",
        ],
        y=1.45,
        h=2.3,
    )
    boxes = row_cards(3, 4.00, 2.2)
    facts = [
        ("Case", "PS000000016 was the posted case."),
        ("Process", "Pid 29320 was killed. A new process started."),
        ("What stayed", "The correlation id stayed. The hold value stayed 0."),
    ]
    bars = [RED, CHAR, GOLD]
    for index, (box, fact) in enumerate(zip(boxes, facts)):
        x, y, w, h = card_shell(slide, box, bars[index])
        add_textbox(
            slide,
            x + 0.22,
            y + 0.28,
            w - 0.4,
            1.6,
            [(fact[0], 15, True, CHAR), (fact[1], 16, False, BODY)],
            name=f"card:{index}",
        )
    speak(
        slide,
        "The same correlation id stayed on PS000000016 after pid 29320 was killed. The hold value stayed 0. This room does not kill a process.",
        "the What stayed card",
        "the five lessons",
    )


def slide_21(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "Five measured results close this walk.")
    lessons = [
        ("1", "Hold prediction scored 117 of 125, the same as the majority guess, and was rejected."),
        ("2", "The lake-built application differed on 2 of 14 checks. 12 checks matched. There is no match percent."),
        ("3", "The correlation id on PS000000016 survived the process restart. The hold value stayed 0."),
        ("4", "Untrusted approve text returned BLOCK. The case record did not change."),
        ("5", "The supporting counts can be shown. Cycle time and dollars stay Unknown. ROI stays educational."),
    ]
    for index, (number, text) in enumerate(lessons):
        y = 1.50 + index * 1.00
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.55), Inches(y + 0.12), Inches(0.48), Inches(0.48))
        fill_solid(circle, RED if index == 0 else CHAR)
        no_line(circle)
        circle.shadow.inherit = False
        write_text(circle, [(number, 16, True, WHITE)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, name="lesson", margin=0.0)
        add_round(slide, 1.25, y, 11.3, 0.85)
        add_textbox(
            slide,
            1.45,
            y,
            10.9,
            0.85,
            [(text, 16, False, CHAR)],
            anchor=MSO_ANCHOR.MIDDLE,
            name="body",
            margin=0.06,
        )
    speak(
        slide,
        "Five measured results close this walk. Read the five sentences. Stop at 40 minutes.",
        "lesson 5",
        "the questions that stay open",
    )


def slide_22(deck: Deck) -> None:
    slide = deck.slide()
    claim(slide, "Ten minutes are open for these questions.")
    add_sentences(
        slide,
        ["These items stay open because the packaged data and the build do not contain an answer. Ask about the page, a count, or one of these limits."],
        y=1.45,
        h=0.9,
    )
    rows = [
        ["Company titles", "RT-01 to RT-06 have role names. Company job titles are Unknown."],
        ["Incident owner", "INC-2025-014 stays empty unless RT-06 types a name on the page."],
        ["Dollar benefit", "ROI stays educational. No cash benefit is claimed."],
        ["Numeric floors", "A floor such as 90 or 95 is Unknown. Freeze 42 of 42 is a count."],
        ["Live model call", "This walk uses replay. A completed Azure call was not run."],
    ]
    draw_table(slide, 0.55, 2.50, [3.0, 9.2], ["Open item", "What is true today"], rows, row_h=0.72)
    speak(
        slide,
        "Ten minutes are open for questions. Company titles, the incident owner, dollars, numeric floors, and a live Azure call stay unanswered by this build.",
        "the open items",
        "Stop. Take the questions.",
    )


def build() -> None:
    deck = Deck()
    for builder in (
        slide_01, slide_02, slide_03, slide_04, slide_05, slide_06,
        slide_07, slide_08, slide_09, slide_10, slide_11, slide_12,
        slide_13, slide_14, slide_15, slide_16, slide_17, slide_18,
        slide_19, slide_20, slide_21, slide_22,
    ):
        builder(deck)
    deck.save()
    scrub_theme(OUT)
    problems = audit(OUT)
    if problems:
        print('AUDIT FAILED')
        for item in problems:
            print('-', item)
        sys.exit(1)
    print(f'Wrote {OUT}')


def scrub_theme(path: Path) -> None:
    theme = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="Fabpulse Defence">
  <a:themeElements>
    <a:clrScheme name="Fabpulse">
      <a:dk1><a:srgbClr val="231F20"/></a:dk1>
      <a:lt1><a:srgbClr val="FFFFFF"/></a:lt1>
      <a:dk2><a:srgbClr val="231F20"/></a:dk2>
      <a:lt2><a:srgbClr val="F2F2F2"/></a:lt2>
      <a:accent1><a:srgbClr val="ED1C24"/></a:accent1>
      <a:accent2><a:srgbClr val="72007D"/></a:accent2>
      <a:accent3><a:srgbClr val="EFBC4A"/></a:accent3>
      <a:accent4><a:srgbClr val="666666"/></a:accent4>
      <a:accent5><a:srgbClr val="4B4B4B"/></a:accent5>
      <a:accent6><a:srgbClr val="FDEBEC"/></a:accent6>
      <a:hlink><a:srgbClr val="231F20"/></a:hlink>
      <a:folHlink><a:srgbClr val="666666"/></a:folHlink>
    </a:clrScheme>
    <a:fontScheme name="Fabpulse Calibri">
      <a:majorFont>
        <a:latin typeface="Calibri"/>
        <a:ea typeface="Calibri"/>
        <a:cs typeface="Calibri"/>
      </a:majorFont>
      <a:minorFont>
        <a:latin typeface="Calibri"/>
        <a:ea typeface="Calibri"/>
        <a:cs typeface="Calibri"/>
      </a:minorFont>
    </a:fontScheme>
    <a:fmtScheme name="Fabpulse">
      <a:fillStyleLst>
        <a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>
        <a:solidFill><a:srgbClr val="F2F2F2"/></a:solidFill>
        <a:solidFill><a:srgbClr val="231F20"/></a:solidFill>
      </a:fillStyleLst>
      <a:lnStyleLst>
        <a:ln w="6350"><a:solidFill><a:srgbClr val="E0E0E0"/></a:solidFill></a:ln>
        <a:ln w="12700"><a:solidFill><a:srgbClr val="231F20"/></a:solidFill></a:ln>
        <a:ln w="19050"><a:solidFill><a:srgbClr val="ED1C24"/></a:solidFill></a:ln>
      </a:lnStyleLst>
      <a:effectStyleLst>
        <a:effectStyle><a:effectLst/></a:effectStyle>
        <a:effectStyle><a:effectLst/></a:effectStyle>
        <a:effectStyle><a:effectLst/></a:effectStyle>
      </a:effectStyleLst>
      <a:bgFillStyleLst>
        <a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>
        <a:solidFill><a:srgbClr val="F2F2F2"/></a:solidFill>
        <a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>
      </a:bgFillStyleLst>
    </a:fmtScheme>
  </a:themeElements>
  <a:objectDefaults/>
  <a:extraClrSchemeLst/>
</a:theme>
"""
    tmp = path.with_suffix(".tmp.pptx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename.startswith("ppt/theme/theme") and info.filename.endswith(".xml"):
                data = theme.encode("utf-8")
            elif info.filename.endswith(".xml"):
                text = data.decode("utf-8")
                for old in ("Aptos", "Aptos Display", "Segoe UI", "Calibri Light", "Carlito"):
                    text = text.replace(old, "Calibri")
                data = text.encode("utf-8")
            target.writestr(info, data)
    tmp.replace(path)


def words(text: str) -> int:
    return len([part for part in text.replace("\n", " ").split(" ") if part.strip()])


def audit(path: Path) -> list[str]:
    problems: list[str] = []
    presentation = Presentation(path)
    if abs(presentation.slide_width - W) > 2000 or abs(presentation.slide_height - H) > 2000:
        problems.append("Slide size is not 13.333 by 7.5 inches.")
    if len(presentation.slides) < 14:
        problems.append("Main arc is shorter than 14 slides.")
    titles = []
    for index, slide in enumerate(presentation.slides, start=1):
        title = ""
        buckets: dict[str, list[str]] = {}
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            name = shape.name or ""
            if name == "title":
                title = text
            if name.startswith("footer") or name == "cover-footer":
                continue
            key = name.split(":")[0]
            buckets.setdefault(key, []).append(text)
            if "birlasoft" in text.lower():
                problems.append(f"Slide {index} contains the word birlasoft.")
        titles.append(title)
        if not title:
            problems.append(f"Slide {index} has no claim title.")
        elif words(title) > 18:
            problems.append(f"Slide {index} title has {words(title)} words: {title}")
        lowered = title.strip().lower().rstrip(".")
        if lowered in {"overview", "agenda", "thank you", "questions", "architecture", "demo", "features", "results"}:
            problems.append(f"Slide {index} title is a topic label: {title}")
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if not notes:
            problems.append(f"Slide {index} has no speaker notes.")
        content_words = 0
        for key, chunks in buckets.items():
            if key == "title":
                continue
            count = words(" ".join(chunks))
            if key == "card" or key.startswith("card"):
                pass
            content_words += count
        # Card caps. Group exact card ids.
        card_groups: dict[str, list[str]] = {}
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            name = shape.name or ""
            if name.startswith("card:"):
                card_groups.setdefault(name, []).append(shape.text_frame.text)
        if index == 1:
            names = {shape.name for shape in slide.shapes}
            if "cover-bar" not in names or "cover-arc-outer" not in names or "cover-arc-inner" not in names:
                problems.append("Cover is missing the bar or an arc.")
        if index > 1:
            for shape in slide.shapes:
                if shape.name == "cover-bar":
                    problems.append(f"Slide {index} repeats the cover bar.")
                if shape.name in {"cover-arc-outer", "cover-arc-inner"}:
                    problems.append(f"Slide {index} repeats a cover arc.")
    print("TITLES")
    for index, title in enumerate(titles, start=1):
        print(f"{index:02d}. {title}")
    with zipfile.ZipFile(path) as package:
        blob = "\n".join(
            package.read(name).decode("utf-8", errors="ignore")
            for name in package.namelist()
            if name.endswith(".xml")
        )
    forbidden = ["143044", "0E7C66", "C9842A", "E8EEF2", "1C2833", "5B9BD5", "Aptos", "Segoe UI", "Calibri Light"]
    for token in forbidden:
        if token.lower() in blob.lower():
            problems.append(f"Forbidden token remains: {token}")
    if "Click to add" in blob:
        problems.append("Placeholder prompt text remains.")
    if 'typeface="Calibri"' not in blob:
        problems.append("Calibri typeface was not written.")
    return problems


if __name__ == "__main__":
    build()
